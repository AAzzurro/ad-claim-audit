"""生成实验报告 Word（目标≥10页）与答辩 PPT。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt

HERE = Path(__file__).resolve().parent


def _set_run_font(run, name_cn: str, name_en: str, size: int, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)


def _set_paragraph_format(p, *, first_indent: bool = False, space_after: float = 6, align="left") -> None:
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if first_indent:
        pf.first_line_indent = Cm(0.74)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 0:
        _set_run_font(run, "黑体", "Times New Roman", 18, True)
        _set_paragraph_format(p, space_after=12, align="center")
    elif level == 1:
        _set_run_font(run, "黑体", "Times New Roman", 14, True)
        _set_paragraph_format(p, space_after=8)
    else:
        _set_run_font(run, "黑体", "Times New Roman", 12, True)
        _set_paragraph_format(p, space_after=6)


def _add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, "宋体", "Times New Roman", 12)
    _set_paragraph_format(p, first_indent=True, space_after=6)


def _add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, "楷体", "Times New Roman", 10.5)
    _set_paragraph_format(p, space_after=4, align="center")


def _shade_header(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): "D9E2F3", qn("w:val"): "clear"})
    tcPr.append(shd)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]], col_cm: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        _set_run_font(run, "黑体", "Times New Roman", 9, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade_header(cell)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            _set_run_font(run, "宋体", "Times New Roman", 9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if col_cm:
        for row in table.rows:
            for i, w in enumerate(col_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def build_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    _add_heading(doc, "短视频广告虚假宣传话术检测系统", 0)
    _add_heading(doc, "实验报告（实践项目三 · 共用稿）", 0)
    meta = doc.add_paragraph()
    run = meta.add_run("作者：【待填】　　学号：【待填】　　日期：2026年9月5日")
    _set_run_font(run, "宋体", "Times New Roman", 12)
    _set_paragraph_format(meta, align="center", space_after=16)

    _add_heading(doc, "摘要", 1)
    _add_body(
        doc,
        "直播带货与短视频广告同时包含口播、字幕、商品卡和促销叠字。审核不能只给“有风险 / 无风险”，"
        "还要留下可复核的原文证据、位置和法规依据。本实验在教师提供的237条约30秒切片上，完成样本信息表、"
        "六类多标签人工标注、SenseVoice口播转写、PP-OCRv6画面识别、关键词规则基线（E1）、"
        "本地Qwen2.5-7B少样本分类（E2），以及规则与模型的标签并集改进（E3）。"
        "风险标签micro F1由E1的0.390升至E2的0.513，再升至E3的0.612；诱导消费F1在融合后达到0.885。"
        "分层复核24条预测，归纳漏检、误检和证据不充分的典型原因。"
        "词表与少样本示例只来自任务书和法规，不把测试视频原句写回规则。"
        "系统以Web Demo「审言」展示视频、ASR/OCR、风险类别、证据和时间戳。",
    )
    _add_body(doc, "关键词：虚假宣传；多标签分类；ASR；OCR；规则基线；少样本学习；证据抽取")

    _add_heading(doc, "1 项目背景与任务", 1)
    _add_body(
        doc,
        "直播间话术常见“行业天花板”“根治”“速效”“无效退款”“稳赚”“保本”“日入”“最后几单”“扫码加群”等表达，"
        "可能构成绝对化、夸大功效、收益承诺、促销误导或站外导流。任务书要求以公开直播回放或短视频广告为输入，"
        "走通“视频处理 → ASR/OCR → 文本风险分类 → 证据抽取 → 效果评价 → Demo”的完整过程，"
        "并借助AI完成编码、提示词和结果分析。本报告按任务书主要任务组织，实验记录对应E1关键词/规则、"
        "E2少样本分类、E3融合改进。",
    )
    _add_body(
        doc,
        "判断原则依据任务书第九条：只依据本条切片中可见、可听到的内容，不推测商品真实效果或商家经营情况；"
        "允许多标签；“可能”“有助于”“效果因人而异”等有条件描述不必然构成风险；"
        "医疗、药品、保健食品、金融投资从严；没有原文证据时应输出“无法判断”。"
        "导向、低俗、地域歧视等原标问题不纳入本任务。",
    )

    _add_heading(doc, "2 样本与标签设计", 1)
    _add_heading(doc, "2.1 样本信息", 2)
    _add_body(
        doc,
        "使用教师提供的公开直播/短视频广告切片，共237条，整包当作测试集。"
        "每条记录sample_id、视频文件名、视频标题、商品/服务类别、来源、采集日期、时长，"
        "见annotations/sample_info.xlsx。时长约28.9–61.4秒，均值约32.5秒。"
        "来源统一记为“教师提供公开直播/短视频广告切片”，采集日期与标注日期对齐为2026-09-04。"
        "视频原文件体积大，提交时可只交转写文本、样本编号和获取说明；本地复现使用data/asr与data/ocr即可。",
    )
    _add_caption(doc, "表1 主要商品/服务类别分布（前8类）")
    _add_table(
        doc,
        ["商品/服务类别", "样本数", "占比"],
        [
            ["0301-普通化妆品", "80", "33.8%"],
            ["0904-厨卫电器", "46", "19.4%"],
            ["0601-保健食品", "23", "9.7%"],
            ["1207-服饰箱包", "16", "6.8%"],
            ["0901-视听设备", "15", "6.3%"],
            ["0510-其他", "10", "4.2%"],
            ["0906-小家电产品", "10", "4.2%"],
            ["1205-清洁用品", "5", "2.1%"],
        ],
        [6.5, 4.0, 4.0],
    )
    _add_body(
        doc,
        "品类以化妆品和家电为主，保健食品与特许加盟数量少，但是法规从严领域。"
        "对疾病治疗用语和收益承诺应采用更严格的标准。样本信息表由抽取索引与人工标注表生成，不改金标准标签。",
    )

    _add_heading(doc, "2.2 标签体系与标注表", 2)
    _add_body(
        doc,
        "标签至少六类，一条可多标。另设“无法判断”处理证据不足。标注表包含ASR口播、字幕/OCR、文案、"
        "风险标签、证据文本、证据时间戳或文本位置、人工说明，见annotations/manual_annotation.xlsx。"
        "原videos/output.xlsx中的导向、低俗等标签仅供参考。推荐方式为规则/模型辅助初标，再按原文人工复核。",
    )
    _add_caption(doc, "表2 金标准标签分布（237条，可多标）")
    _add_table(
        doc,
        ["标签", "口径摘要", "条数"],
        [
            ["正常", "普通介绍、常规促销、有条件描述", "32"],
            ["存在夸大功效", "功效断言、医疗用语、绝对化描述", "195"],
            ["存在虚假收益承诺", "稳赚、保本、日入、保过等", "1"],
            ["存在诱导消费", "限时稀缺、虚构最低价", "27"],
            ["存在站外导流风险线索", "扫码加群、加微信、私信下单", "0"],
            ["存在其他线索", "四类未覆盖，如迷信好运", "1"],
            ["无法判断", "原文过短或证据不足", "1"],
        ],
        [5.2, 7.0, 2.2],
    )
    _add_body(
        doc,
        "多标签20条，其中“夸大功效+诱导消费”19条，“夸大功效+虚假收益承诺”1条（样本192）。"
        "站外导流在本包金标准中支持为0，规则和提示词仍保留该类，以便Demo处理新视频。"
        "绝对化描述（如行业天花板、全国第一）按细则归入夸大功效，而不是单独开“绝对化”标签。",
    )

    _add_heading(doc, "3 视频预处理与文字信息获取", 1)
    _add_heading(doc, "3.1 音频与口播转写", 2)
    _add_body(
        doc,
        "src/media.py使用ffmpeg抽取wav。默认ASR为FunASR SenseVoiceSmall，输出带时间戳的分段和全文，"
        "并去掉SenseVoice的语言与情感标记。备选后端为Whisper。全量237条抽取成功，写入data/asr/{id}.json。"
        "口播质量受直播噪音、方言和顺口溜影响。部分切片口播几乎为空（如样本210），此时必须以画面为准，"
        "或按任务书9(5)标“无法判断”。",
    )
    _add_heading(doc, "3.2 抽帧与OCR", 2)
    _add_body(
        doc,
        "每隔3秒抽取一帧，相邻帧画面变化小于阈值则跳过，避免对静止商品卡重复识别。"
        "OCR使用PP-OCRv6 small检测+识别。src/merge.py去掉公屏、进场提示、控件和明显重复行，"
        "保留右侧商品卡与卖点叠字。结果写入data/ocr与data/merged。"
        "虚假宣传话术经常只出现在画面：“天花板”“全年最低价”“活动就一天”“国内首个”都可能完全不入口播。"
        "因此后续分类必须同时读取ASR与OCR。OCR仍会把“无毒无害”认成“无海无害”，叠字重复也会干扰模型。",
    )
    _add_heading(doc, "3.3 清洗原则", 2)
    _add_body(
        doc,
        "清洗只做去重和控件过滤，不根据测试集词频删除“疑似风险”句子，以免评测泄漏。"
        "合并后的audio_text与visual_text作为E1、E2的共同输入，保证方法对比时文本条件一致。",
    )

    _add_heading(doc, "4 规则基线（E1）", 1)
    _add_body(
        doc,
        "E1是小规模关键词与正则基线，实现于src/rules.py与src/detect.py。"
        "词表只允许从Knowledge/写入：任务书示例（行业天花板、根治、速效、无效退款、稳赚、保本、日入、"
        "最后几单、扫码加群）以及《广告法》第九、十六至十八、二十四、二十五条和《直播电商监督管理办法》第三十二条。"
        "禁止根据本包句子或P/R/F1漏检加词。这是本实验的硬约束，也是E3选择融合而不是扩词表的原因。",
    )
    _add_body(
        doc,
        "规则分为五组。绝对化归入夸大功效，包括国家级、最高级、最佳、第一、天花板等，但不用单字“最”，"
        "以免“最高领取”误伤。功效/医疗包括根治、治愈、速效、三天见效、无效退款、治疗功能。"
        "收益承诺包括稳赚、保本、日入、包过。诱导消费包括最后几单/一批/一天以及全网/全年/历史最低价。"
        "站外导流包括扫码加群、加微信、私信下单。命中点附近若出现“可能”“有助于”“效果因人而异”，"
        "按任务书9(3)改为无法判断。无命中则输出正常，并在解释中写明“未命中不等于人工认定无风险”。",
    )
    _add_body(
        doc,
        "E1在237条上的风险标签micro精确率为0.948，召回仅0.246，micro F1为0.390，完全匹配72/237。"
        "诱导消费F1为0.773，夸大功效F1为0.322。虚假收益承诺与其他线索均为0，"
        "因为词表没有“只谈收益”“带来好运”这类写法。规则的价值是证据可回原文、依据可写到具体条款，"
        "适合作为审核台的可解释底线。",
    )

    _add_heading(doc, "5 语言模型少样本分类（E2）", 1)
    _add_heading(doc, "5.1 模型与提示词", 2)
    _add_body(
        doc,
        "任务书要求使用小规模LM（示例Qwen8B同档），不得直接使用28B。本实验使用本机Ollama部署的qwen2.5:7b。"
        "系统提示要求：只根据本条ASR/OCR原文判断；先读画面叠字再读口播；导向低俗不属于本任务；"
        "每一项风险必须带能在原文落地的短证据；医疗保健金融从严。"
        "同时明确列出不要标成风险的直播套话：最高领取N元、拍下链接、粉丝真实反馈。",
    )
    _add_body(
        doc,
        "少样本共8组，全部来自任务书背景句和法规表述，例如“行业天花板/根治/三天见效”“稳赚不赔”"
        "“最后几单/全网最低价”“扫码加群”“带来好运”“可能有助于/效果因人而异”，以及空文本对应无法判断。"
        "没有使用videos/中的任何原句。这保证E2的提升来自泛化，而不是背测试集。",
    )
    _add_heading(doc, "5.2 结构化输出与后处理", 2)
    _add_body(
        doc,
        "模型只允许输出一个JSON对象，字段与E1对齐。src/classify.py随后做标签白名单、"
        "证据接地（证据必须是口播或OCR的子串）、风险与“正常/无法判断”互斥。"
        "解释里若提到站外导流但证据不合格，标签会被丢掉。该后处理把风险标签精确率维持在0.88左右，"
        "明显减少把领券、拍链接当成风险的情况。",
    )
    _add_heading(doc, "5.3 E2结果", 2)
    _add_body(
        doc,
        "237条全部出数。风险micro P/R/F1为0.880/0.362/0.513，macro F1为0.423，完全匹配89/237（37.6%）。"
        "夸大功效F1由0.322升到0.495，能够抓住“肌肤老10倍”“14天改善”等词表外断言；"
        "诱导消费F1由0.773降到0.619，常把“最低价”当成常规促销；其他线索1条命中；收益承诺仍漏。"
        "与E1对照，两者精确率都高、召回都低，且强项不同，具备融合条件。",
    )

    _add_heading(doc, "6 输出格式", 1)
    _add_body(
        doc,
        "E1、E2、E3与Demo共用同一套字段，便于人工对照。risk_labels为多标签列表；"
        "evidence为原文短句；evidence_position给出source（audio/visual）、time和span；"
        "rule_basis写任务书位置或法条；explanation用一两句话说明依据哪句原文。"
        "Demo另有verdict（risk/normal/unknown）。模型不得推测商品是否真有效，只能引用看得见、听得见的原文。"
        "若证据不足，应输出无法判断，而不是用常识补全。",
    )

    _add_heading(doc, "7 方法对比结果", 1)
    _add_body(
        doc,
        "测试集237条，金标准见annotations/manual_annotation.json。"
        "指标按风险五类做多标签Precision、Recall和F1，同时报告完全匹配率。"
        "表3对应任务书建议的最小实验记录。E3细节见第10节。",
    )
    _add_caption(doc, "表3 最小实验记录（风险标签micro指标）")
    _add_table(
        doc,
        ["实验", "方法/条件", "P", "R", "F1", "完全匹配", "结论"],
        [
            ["E1", "关键词/规则基线", "0.948", "0.246", "0.390", "72/237", "精确高，召回低"],
            ["E2", "Qwen2.5-7B少样本", "0.880", "0.362", "0.513", "89/237", "功效补上，促销变弱"],
            ["E3", "E1∪E2标签并集", "0.897", "0.464", "0.612", "106/237", "不改词表即可提升"],
        ],
        [1.6, 3.6, 1.5, 1.5, 1.5, 2.2, 3.2],
    )
    _add_caption(doc, "表4 分标签F1对比")
    _add_table(
        doc,
        ["标签", "支持", "E1 F1", "E2 F1", "E3 F1", "E3 P", "E3 R"],
        [
            ["存在夸大功效", "195", "0.322", "0.495", "0.561", "0.889", "0.410"],
            ["存在虚假收益承诺", "1", "0.000", "0.000", "0.000", "0.000", "0.000"],
            ["存在诱导消费", "27", "0.773", "0.619", "0.885", "0.920", "0.852"],
            ["存在站外导流风险线索", "0", "—", "—", "—", "—", "—"],
            ["存在其他线索", "1", "0.000", "1.000", "1.000", "1.000", "1.000"],
        ],
        [4.0, 1.6, 1.7, 1.7, 1.7, 1.6, 1.6],
    )
    _add_body(
        doc,
        "E3没有引入测试集新词，只是把已经算过的预测并起来。诱导消费召回从E2的0.482升到0.852，"
        "接近并超过单独E1；夸大功效召回从E2的0.344升到0.410。精确率仍保持约0.90，"
        "说明并集没有把大量误检带进来——原因是两路本身精确率都高、召回都低。"
        "收益承诺（192）两边都漏，融合无法挽回。无法判断（210）在融合规则里被E1的“正常”盖掉，"
        "这是并集策略的代价。",
    )

    _add_heading(doc, "8 人工复核结果", 1)
    _add_body(
        doc,
        "按任务书要求，从E1/E2分歧和漏检中分层抽取24条（多于20条），对照金标准与ASR/OCR原文填写复核表。"
        "字段包括样本编号、人工标签、E1/E2模型标签、证据是否正确、是否需要修改、误判原因，"
        "见annotations/review_20.xlsx。复核阶段不改词表或提示词，符合“先分析、再改进”的实验顺序。",
    )
    _add_body(
        doc,
        "抽样构成：稀有标签191/192/210；已知问题10、162；E1/E2互补20、18、78、129、150；"
        "双边漏检6、16、24、94；误检17、136、167、208、216；完全匹配正负例1、7。"
        "摘要上，E1有16条需要修改，E2同样16条；E1证据正确或部分正确8条，E2为10条。"
        "两边改动数量接近，但错误类型不同，说明不是同一套失败模式。",
    )
    _add_caption(doc, "表5 复核中归纳的主要误判类型")
    _add_table(
        doc,
        ["类型", "代表样本", "主要表现"],
        [
            ["词表外功效漏检", "6, 16, 24", "即刻紧致、无毒无害、瓦解黑色素"],
            ["价格误导被模型放过", "18, 44", "最低价被当成常规促销"],
            ["天花板字面误伤", "208, 216", "预算上限或“没有天花板”"],
            ["标签对证据错", "162", "应引治疗骨关节炎，却取可达0.4"],
            ["稀有标签", "191, 192", "好运E2对；收益承诺两边漏"],
            ["催拍无关键词", "94", "倒计时321不含最后几单"],
            ["证据不足", "210", "E2标无法判断，E1标正常"],
        ],
        [3.8, 3.2, 8.0],
    )
    _add_body(
        doc,
        "这些观察直接支持第10节的改进选择：用并集吃互补，而不是根据漏检把本包句子写进词表。"
        "若在复核之后把“治疗骨关节炎”“不谈模式只谈收益”写进规则，表面上F1会涨，但违反测试集纪律，"
        "也不能证明方法对未见视频有效。",
    )

    _add_heading(doc, "9 典型正确与错误案例", 1)
    _add_heading(doc, "9.1 正确正例：样本7", 2)
    _add_body(
        doc,
        "海信激光电视画面同时出现“行业唯一”和“色域天花板”。E1命中天花板，E2两条都引用。"
        "标签与金标准一致。说明任务书示例词出现在画面上时，规则和模型都可以稳定工作，"
        "此时审核解释也能写到《广告法》第九条第三项。",
    )
    _add_heading(doc, "9.2 互补漏检：样本10与20", 2)
    _add_body(
        doc,
        "样本10金标准为夸大功效+诱导消费。E1只打中口播“天花板”，E2只打中画面“活动就一天”。"
        "样本20类似：E1打中“水蜜桃片的也个天花板”，E2打中“最后十四箱”。"
        "各自证据都能在原文落地，单路标签却不完整。并集后与金标准对齐。这是E3最直接的动机，"
        "也说明Demo提供“规则+模型”模式不是界面点缀，而是评测上站得住的策略。",
    )
    _add_heading(doc, "9.3 标签对、证据错：样本162", 2)
    _add_body(
        doc,
        "保健食品切片口播出现“盐酸氨糖治疗骨关节炎”，同时口头声明“不涉及治疗”。"
        "按任务书从严和高敏感领域条款，仍应标夸大功效，证据应落在“治疗骨关节炎”。"
        "E1无单独“治疗”规则，输出正常。E2标签正确，证据却取了效应量“可达0.4”。"
        "复核结论是需要修改证据，而不是改金标准。后续若从《广告法》第十七条补“治疗”二字，"
        "属于法规同义，仍然不得写入本条原句。",
    )
    _add_heading(doc, "9.4 双边漏检：样本192", 2)
    _add_body(
        doc,
        "加盟广告画面有“国内首个线上AI测肤功能”和“不谈模式只谈收益”。"
        "合并文本里这两句都在，并不是抽取失败。E1词表没有“首个”，也没有“只谈收益”；"
        "E2在较长OCR中输出正常。这是本包唯一的虚假收益承诺，E1/E2/E3该类F1均为0。"
        "报告把它列为当前召回上限，用来说明稀有类不能只看一个F1数字。",
    )
    _add_heading(doc, "9.5 误检：样本208与216", 2)
    _add_body(
        doc,
        "样本208“学生党的耳机价格天花板也就几百块”是预算上限，E1误伤，E2正确放行。"
        "样本216“吃梅才真的没有天花板”是程度修辞，两边都误伤。"
        "说明“天花板”规则过宽，需要否定或预算语境，但不能用本包这两句当新正则的唯一来源。",
    )
    _add_heading(doc, "9.6 正确负例：样本1", 2)
    _add_body(
        doc,
        "口播属于导向/性别对立，画面有“最高领取20元”。E1与E2都标正常，"
        "符合“本任务只标虚假宣传、领券不是最高级”。这条用来检查系统有没有把原数据集的其它审核口径带进来。",
    )

    _add_heading(doc, "10 一次简单改进及验证（E3）", 1)
    _add_body(
        doc,
        "任务书允许的改进方向包括增加同义规则、改进文本清洗、加入少样本、法规条款检索、优化提示词等。"
        "受词表纪律约束，本实验选择不改词表、不改提示词、不重跑237次LLM，"
        "只把已有E1与E2的风险标签取并集，证据去重后一并输出。"
        "融合逻辑与Demo的“规则+模型”模式相同，批量入口为python -m src.hybrid，评测为experiment=e3。",
    )
    _add_body(
        doc,
        "改进前（E2）micro F1为0.513，诱导消费F1为0.619，夸大功效F1为0.495。"
        "改进后（E3）分别为0.612、0.885、0.561。完全匹配由89升到106。"
        "诱导消费漏检从14降到4，夸大功效真阳性从67升到80。"
        "精确率仍保持0.90左右。该改进满足任务书“比较改进前后的指标或典型案例效果”，"
        "并且可以在没有Ollama的机器上复现，因为只读data/detect与data/classify。",
    )
    _add_body(
        doc,
        "未解决的问题仍然清楚：192收益承诺、大量词表外功效（夸大功效仍有115条漏检）、"
        "210被融合成正常。后续方向均不得使用测试集原句：从法规补“治疗”“首个/首创”等同义；"
        "给“天花板”加否定窗；空口播且无画面命中时优先无法判断；"
        "对保健食品提高疾病治疗用语的权重。这些方向写进报告，本轮不落地到词表。",
    )

    _add_heading(doc, "11 审核结果展示Demo", 1)
    _add_body(
        doc,
        "编写了简单Web页面，满足任务书第7项。执行python -m src.app后，默认打开http://127.0.0.1:7860。"
        "页面名称为「审言」。使用者可以上传mp4，或从237条测试集中选择样本。"
        "三种分析模式分别为规则基线、模型分类、规则+模型。"
        "页面展示视频基本信息、口播分段文本、OCR叠字、风险类别、原文证据、时间戳、规则依据和审核解释。"
        "不要求账号登录、在线爬取或自动处罚。已抽取样本会复用缓存；新视频会现场跑SenseVoice与OCR。"
        "模型模式需要本机Ollama可用。命令行等价入口为src.detect、src.classify与src.hybrid。",
    )

    _add_heading(doc, "11.1 AI协作如何进入本实验", 2)
    _add_body(
        doc,
        "任务书要求记录至少两条关键AI使用情况。第一条用于设计结构化输出和少样本约束："
        "向AI询问如何让7B模型给出可复核JSON，并避免把拍链接、领券标成风险。"
        "建议被采用，落实为src/prompts.py与classify.py的证据接地。"
        "第二条用于选择改进策略：在不能把测试集句子写进词表的前提下，AI建议先复核再做E1∪E2融合。"
        "该建议被采用，并得到第7节的指标支持。完整问答见reports/AI协作记录.md。",
    )

    _add_heading(doc, "12 结论与分工说明", 1)
    _add_body(
        doc,
        "本实践完成了任务书规定的样本准备、标注、ASR/OCR、规则基线、小规模LM分类、指标评估、"
        "不少于20条人工复核、一次简单改进和Demo。主要结论有五条。"
        "第一，虚假宣传话术大量出现在画面叠字，只做ASR不够。"
        "第二，冻结词表的规则精确、召回窄；少样本模型能补词表外功效，但促销误导不稳。"
        "第三，在禁止测试集泄漏的前提下，标签并集是有效且可复现的简单改进。"
        "第四，稀有类支持不足，复核比单一F1更能说明系统边界。"
        "第五，证据质量与标签同等重要，162说明标签对、证据错仍不合格。",
    )
    _add_body(
        doc,
        "书面材料另附AI协作记录、小组分工说明和答辩PPT。组员姓名在提交前填入【待填】。"
        "建议三人分组时，一人负责样本与标注，一人负责抽取与规则，一人负责模型、评测与报告；"
        "四人时再拆出Demo与答辩。各人提交的报告可在本共用稿基础上增加个人负责模块的细节。",
    )

    _add_heading(doc, "参考文献与材料", 1)
    for item in (
        "[1] 实践项目任务书（三）：短视频广告虚假宣传话术检测。",
        "[2] 《中华人民共和国广告法》第九条、第十六条至第十八条、第二十四条、第二十五条。",
        "[3] 《直播电商监督管理办法》第三十二条、第三十四条。",
        "[4] 本仓库源代码src/、标注annotations/、评测data/eval/与书面材料reports/。",
    ):
        p = doc.add_paragraph()
        run = p.add_run(item)
        _set_run_font(run, "宋体", "Times New Roman", 12)
        _set_paragraph_format(p, space_after=4)

    _add_heading(doc, "附录A 复现命令", 1)
    _add_body(
        doc,
        "环境安装执行./setup_env.sh后激活.venv。评测E1：python -m src.detect 再 python -m src.eval。"
        "评测E2需本机Ollama：python -m src.classify，再指定--detect-dir data/classify --experiment e2。"
        "E3：python -m src.hybrid 与 python -m src.eval --detect-dir data/hybrid --experiment e3。"
        "Demo：python -m src.app。无模型机器可跳过抽取与E2，直接评已提交的预测目录。",
    )
    _add_heading(doc, "附录B 复核样本编号", 1)
    _add_body(
        doc,
        "1、6、7、10、14、16、17、18、20、24、44、50、78、94、129、136、150、162、167、191、192、208、210、216。"
        "完整字段见annotations/review_20.xlsx。组员抽查时建议至少重看10、162、192、208四条。",
    )
    _add_caption(doc, "表6 复核表字段与统计")
    _add_table(
        doc,
        ["项目", "数量", "说明"],
        [
            ["复核条数", "24", "多于任务书要求的20条"],
            ["E1需要修改", "16", "漏检为主，另有天花板误伤"],
            ["E2需要修改", "16", "漏最低价、误标卖点、证据取偏"],
            ["E1证据正确/部分正确", "8", "命中后span一般可回原文"],
            ["E2证据正确/部分正确", "10", "存在重复堆叠和取错短句"],
        ],
        [5.0, 3.0, 7.0],
    )

    _add_heading(doc, "附录C 系统流程说明", 1)
    _add_body(
        doc,
        "视频进入后先探测时长并抽音、抽帧。ASR与OCR可以跳过已有缓存。"
        "规则引擎在口播分段和OCR合并行上做正则命中，记录span与时间戳。"
        "语言模型接收拼接后的【口播ASR】与【画面OCR】，返回JSON，再做接地过滤。"
        "融合器只在风险标签集合上取并，解释拼接“规则命中/模型命中/谁补上了哪一类”。"
        "评测脚本不回写规则。Web层用SSE推送加载模型、抽取、检测、融合等阶段，便于课堂演示。",
    )
    _add_body(
        doc,
        "若设备无法运行ASR/OCR，任务书允许提交已提取文本和清晰复现说明。"
        "本仓库已包含237条转写与识别结果，以及E1/E2/E3预测和指标文件。"
        "因此评阅人可以在不下载视频、不下载本地权重的情况下复核数字和案例。",
    )

    out = HERE / "实验报告.docx"
    doc.save(out)
    return out


NAVY = PptRGB(0x1F, 0x3A, 0x5F)
TEAL = PptRGB(0x2A, 0x6F, 0x6F)
INK = PptRGB(0x22, 0x22, 0x22)
MUTED = PptRGB(0x55, 0x55, 0x55)
WHITE = PptRGB(0xFF, 0xFF, 0xFF)


def _slide_box(slide, left, top, width, height, text, *, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "PingFang SC"
    return tf


def _add_bullets(tf, lines: list[str], size=18) -> None:
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 and not tf.paragraphs[0].runs else tf.add_paragraph()
        if i == 0 and tf.paragraphs[0].runs:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        run.font.size = PptPt(size)
        run.font.color.rgb = INK
        run.font.name = "PingFang SC"


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PptRGB(0xF7, 0xF4, 0xEE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    return slide


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = _blank_slide(prs)
    _slide_box(s, 0.8, 2.0, 11.5, 1.0, "短视频广告虚假宣传话术检测", size=36, bold=True, color=NAVY)
    _slide_box(s, 0.8, 3.2, 11.5, 0.6, "实践项目（三）　答辩材料　共用稿", size=20, color=TEAL)
    _slide_box(s, 0.8, 5.8, 11.5, 0.5, "作者：【待填】　　2026-09-05", size=16, color=MUTED)

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "任务与目标", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "输入：约30秒直播/短视频广告切片（本包237条，整包作测试集）",
            "过程：抽音抽帧 → ASR/OCR → 规则/LM分类 → 证据与法条 → 评测 → Demo",
            "输出：风险标签、原文证据、时间戳/位置、规则依据、解释",
            "约束：小规模LM（Qwen2.5-7B）；词表与少样本不得使用测试视频原句",
            "标签：正常 / 夸大功效 / 虚假收益 / 诱导消费 / 站外导流 / 其他线索 / 无法判断",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "样本与金标准", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "237条，时长均值32.5秒；化妆品80、厨卫电器46、保健食品23",
            "夸大功效195，诱导消费27，正常32，多标签20",
            "收益承诺、其他线索、无法判断各1条；站外导流支持为0",
            "判断只看本条可见可听内容，不推测商品是否真有效",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "ASR / OCR", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "ASR：FunASR SenseVoiceSmall，带时间戳，237/237成功",
            "OCR：每3秒抽帧 + 画面变化去重，PP-OCRv6 small",
            "清洗：去公屏/控件，保留商品卡与卖点叠字",
            "关键观察：最低价、天花板、首个，经常只出现在画面",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "E1 规则基线", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "词表只来自任务书与《广告法》《直播电商监督管理办法》",
            "micro P/R/F1 = 0.948 / 0.246 / 0.390，完全匹配72/237",
            "诱导消费稳（F1 0.773），夸大功效召回低（F1 0.322）",
            "有条件描述附近命中改为无法判断；无命中≠人工认定安全",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "E2 少样本分类", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "本机 Ollama Qwen2.5-7B；8组少样本均来自任务书/法规",
            "后处理：证据必须是原文子串，否则丢标签",
            "micro P/R/F1 = 0.880 / 0.362 / 0.513，完全匹配89/237",
            "能抓“老10倍”“14天改善”；常把“最低价”放过",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "E3 简单改进：规则 ∪ 模型", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "不改词表、不改提示词、不重跑LLM，只融合已有预测",
            "micro F1：0.390 → 0.513 → 0.612",
            "诱导消费 F1：0.773 / 0.619 → 0.885（召回0.852）",
            "夸大功效 F1：0.322 / 0.495 → 0.561；完全匹配106/237",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "人工复核（24条）", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "分层：稀有标签、互补、漏检、误检、正负例",
            "E1、E2各有16条需要修改，但错因不同",
            "典型：10/20互补；162证据错；192收益承诺双边漏",
            "208预算天花板误伤；94倒计时催拍无关键词",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "案例：互补与证据", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "10：E1打中天花板，E2打中“活动就一天”，并集才完整",
            "162：应引“治疗骨关节炎”，模型却取“可达0.4”",
            "192：画面有“只谈收益/国内首个”，词表与模型都漏",
            "1：导向口播+最高领取，正确标正常",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "Demo「审言」", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "python -m src.app  →  http://127.0.0.1:7860",
            "上传视频或从测试集选择；三种模式可切换",
            "展示口播、OCR、标签、证据、时间戳、法条、解释",
            "无登录、无爬取、无自动处罚",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 0.4, 12, 0.6, "结论", size=28, bold=True, color=NAVY)
    tf = _slide_box(s, 0.8, 1.3, 11.5, 5.5, "", size=18)
    _add_bullets(
        tf,
        [
            "画面叠字与口播必须一起审",
            "规则保精确和可解释，模型补词表外断言",
            "纪律内的并集是有效的一次改进",
            "证据质量与标签同等重要；稀有类要靠复核说话",
        ],
        20,
    )

    s = _blank_slide(prs)
    _slide_box(s, 0.7, 2.4, 12, 1.0, "谢谢。欢迎提问。", size=36, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _slide_box(s, 0.7, 3.6, 12, 0.5, "姓名处请改为组员名单", size=16, color=MUTED, align=PP_ALIGN.CENTER)

    out = HERE / "答辩.pptx"
    prs.save(out)
    return out


def main() -> None:
    docx = build_docx()
    pptx = build_pptx()
    print("wrote", docx)
    print("wrote", pptx)


if __name__ == "__main__":
    main()
